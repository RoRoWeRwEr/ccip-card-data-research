#!/usr/bin/env python3
"""Add CC77 staging/provenance sheets while preserving every existing sheet."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs" / "excel" / "saudi-credit-cards-unified-consolidated.xlsx"
MACHINE = ROOT / "outputs" / "machine-readable"
SHEETS = ["CC77 Source Registry", "CC77 Comparison Leads", "CC77 VAT Review"]
NAVY = "1F4E78"
PALE_BLUE = "D9EAF7"
PALE_GOLD = "FFF2CC"
PALE_GREEN = "E2F0D9"


def load_json(name: str):
    return json.loads((MACHINE / name).read_text(encoding="utf-8"))


def header_style(ws, row: int, columns: int) -> None:
    for cell in ws.iter_cols(min_row=row, max_row=row, min_col=1, max_col=columns):
        target = cell[0]
        target.fill = PatternFill("solid", fgColor=NAVY)
        target.font = Font(color="FFFFFF", bold=True)
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def body_style(ws, start_row: int, end_row: int, columns: int, height: int = 54) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=columns):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for index in range(start_row, end_row + 1):
        ws.row_dimensions[index].height = height


def add_table(ws, ref: str, name: str, style: str = "TableStyleMedium2") -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=style, showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def main() -> None:
    summary = load_json("cc77_source_summary.json")
    links = load_json("cc77_embedded_links.json")
    leads = load_json("cc77_leads.json")
    vat_rows = load_json("cc77_vat_review.json")

    wb = load_workbook(WORKBOOK)
    for name in SHEETS:
        if name in wb.sheetnames:
            raise SystemExit(f"{name} already exists; rebuild before rerunning this additive script")

    source = wb.create_sheet(SHEETS[0])
    source.sheet_view.showGridLines = False
    source.freeze_panes = "A19"
    source.append(["Source metadata", "Value", "", "Control metric", "Value"])
    metadata = [
        ("Source ID", summary["source_id"]),
        ("Repository path", summary["path"]),
        ("SHA-256", summary["sha256"]),
        ("File size (bytes)", summary["file_size_bytes"]),
        ("Pages", summary["pages"]),
        ("Displayed version", summary["displayed_version"]),
        ("Displayed update date", summary["displayed_update_date"]),
        ("Source type", summary["source_type"]),
        ("Authority rank", summary["authority_rank"]),
        ("Publisher identity", summary["publisher_identity"]),
        ("Embedded annotations", summary["embedded_annotations"]),
        ("Embedded supporting links", summary["embedded_supporting_links"]),
        ("Extraction date", summary["extraction_date"]),
        ("Existing master data changed", "no"),
        ("Safety decision", "Secondary lead source only; no overwrite, merge, deletion, ID assignment, or confidence promotion."),
    ]
    controls = [
        ("Embedded supporting links", "=COUNTA(A19:A31)"),
        ("Staged comparison leads", "=COUNTA('CC77 Comparison Leads'!A2:A14)"),
        ("VAT review records", "=COUNTA('CC77 VAT Review'!A2:A6)"),
        ("Existing master rows overwritten", "=0"),
        ("Existing master rows deleted", "=0"),
    ]
    for index, (key, value) in enumerate(metadata, 2):
        source.cell(index, 1, key)
        source.cell(index, 2, value)
    for index, (key, value) in enumerate(controls, 2):
        source.cell(index, 4, key)
        source.cell(index, 5, value)
    for cell in source[1]:
        if cell.column != 3:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(2, 17):
        source.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        source.cell(row, 1).font = Font(bold=True)
        source.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
        source.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(2, 7):
        source.cell(row, 4).fill = PatternFill("solid", fgColor=PALE_GREEN)
        source.cell(row, 4).font = Font(bold=True)
        source.cell(row, 4).alignment = Alignment(vertical="top", wrap_text=True)
        source.cell(row, 5).alignment = Alignment(vertical="top", wrap_text=True)
    link_headers = list(links[0])
    for column, header in enumerate(link_headers, 1):
        source.cell(18, column, header)
    for row_index, row in enumerate(links, 19):
        for column, header in enumerate(link_headers, 1):
            source.cell(row_index, column, row[header])
    header_style(source, 18, len(link_headers))
    body_style(source, 19, 18 + len(links), len(link_headers), 60)
    add_table(source, f"A18:L{18 + len(links)}", "CC77EmbeddedLinks")
    widths = {"A": 24, "B": 58, "C": 34, "D": 54, "E": 48, "F": 18, "G": 22, "H": 34, "I": 20, "J": 20, "K": 20, "L": 58}
    for column, width in widths.items():
        source.column_dimensions[column].width = width

    lead_sheet = wb.create_sheet(SHEETS[1])
    lead_sheet.sheet_view.showGridLines = False
    lead_sheet.freeze_panes = "A2"
    lead_headers = list(leads[0])
    lead_sheet.append(lead_headers)
    for row in leads:
        lead_sheet.append([row[header] for header in lead_headers])
    header_style(lead_sheet, 1, len(lead_headers))
    body_style(lead_sheet, 2, 1 + len(leads), len(lead_headers), 72)
    add_table(lead_sheet, f"A1:I{1 + len(leads)}", "CC77ComparisonLeads")
    for column, width in zip("ABCDEFGHI", [16, 28, 36, 32, 42, 68, 52, 28, 72]):
        lead_sheet.column_dimensions[column].width = width

    vat_sheet = wb.create_sheet(SHEETS[2])
    vat_sheet.sheet_view.showGridLines = False
    vat_sheet.freeze_panes = "A2"
    vat_headers = list(vat_rows[0])
    vat_sheet.append(vat_headers)
    for row in vat_rows:
        vat_sheet.append([row[header] for header in vat_headers])
    header_style(vat_sheet, 1, len(vat_headers))
    body_style(vat_sheet, 2, 1 + len(vat_rows), len(vat_headers), 72)
    for row in vat_sheet.iter_rows(min_row=2, max_row=1 + len(vat_rows), min_col=1, max_col=len(vat_headers)):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PALE_GOLD)
    add_table(vat_sheet, f"A1:G{1 + len(vat_rows)}", "CC77VATReview", "TableStyleMedium9")
    for column, width in zip("ABCDEFG", [18, 34, 48, 48, 38, 36, 72]):
        vat_sheet.column_dimensions[column].width = width

    temp = WORKBOOK.with_suffix(".cc77.tmp.xlsx")
    wb.save(temp)
    temp.replace(WORKBOOK)
    print(json.dumps({"workbook": str(WORKBOOK), "added_sheets": SHEETS, "links": len(links), "leads": len(leads), "vat_rows": len(vat_rows)}, indent=2))


if __name__ == "__main__":
    main()
